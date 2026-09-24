"""
Risk of Bias (RoB) & Study Quality Assessment Matrix Template
-------------------------------------------------------------
Generates publication-standard Risk of Bias assessment tables (ROBINS-I / Cochrane
/ Newcastle-Ottawa Scale criteria) with OpenXML color-coded badge cells.
"""

import docx
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_bg(cell, fill_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_pad(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_rob_matrix_table(doc, studies_data, header_bg="1F4E79"):
    """
    Renders a publication-grade Risk of Bias matrix table with color-coded risk ratings.
    
    studies_data: list of dicts:
    [
        {
            'study': 'Plastina et al. (2020)',
            'design': 'Producer Survey (n=1,200)',
            'selection': 'Low',
            'confounding': 'Low',
            'measurement': 'Low',
            'attrition': 'Moderate',
            'reporting': 'Low',
            'overall': 'Low'
        },
        ...
    ]
    """
    headers = [
        "Study Citation",
        "Empirical Design",
        "Selection Bias",
        "Confounding",
        "Measurement",
        "Attrition",
        "Reporting",
        "Overall RoB"
    ]
    
    tbl = doc.add_table(rows=len(studies_data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Format Header Row
    for col_idx, h_text in enumerate(headers):
        cell = tbl.cell(0, col_idx)
        set_cell_bg(cell, header_bg)
        set_cell_pad(cell, top=80, bottom=80, left=90, right=90)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.name = 'Calibri'
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Risk badge color map (Soft background / Deep text)
    risk_styles = {
        'low': ('E8F8F5', RGBColor(20, 90, 50), '[+] Low'),
        'moderate': ('FEF9E7', RGBColor(125, 102, 8), '[~] Moderate'),
        'high': ('FDEDEC', RGBColor(146, 43, 33), '[-] High')
    }

    # Data Rows
    for row_idx, item in enumerate(studies_data, start=1):
        # 1. Study Citation
        c0 = tbl.cell(row_idx, 0)
        set_cell_bg(c0, "FFFFFF" if row_idx % 2 == 1 else "F9FBFD")
        set_cell_pad(c0, 60, 60, 70, 70)
        r0 = c0.paragraphs[0].add_run(item['study'])
        r0.font.name = 'Calibri'
        r0.font.bold = True
        r0.font.size = Pt(8.5)

        # 2. Design
        c1 = tbl.cell(row_idx, 1)
        set_cell_bg(c1, "FFFFFF" if row_idx % 2 == 1 else "F9FBFD")
        set_cell_pad(c1, 60, 60, 70, 70)
        r1 = c1.paragraphs[0].add_run(item['design'])
        r1.font.name = 'Calibri'
        r1.font.size = Pt(8.5)

        # Domains (cols 2 to 7)
        domain_keys = ['selection', 'confounding', 'measurement', 'attrition', 'reporting', 'overall']
        for d_idx, d_key in enumerate(domain_keys, start=2):
            val = str(item.get(d_key, 'Low')).strip().lower()
            cell = tbl.cell(row_idx, d_idx)
            bg, text_color, label = risk_styles.get(val, ('F2F4F7', RGBColor(80, 90, 100), val.capitalize()))
            
            set_cell_bg(cell, bg)
            set_cell_pad(cell, 60, 60, 60, 60)
            p = cell.paragraphs[0]
            r = p.add_run(label)
            r.font.name = 'Calibri'
            r.font.bold = True
            r.font.size = Pt(8)
            r.font.color.rgb = text_color

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(6)
    return tbl
